
from openpyxl import Workbook, load_workbook
import os

def write_excel(filename):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Name", "Age"])
    sheet.append(["john",30])
    sheet.append(["smith",35])
    workbook.save(filename)


def read_excel(filename):
    workbook  = load_workbook(filename)
    sheet = workbook.active
    for now in sheet.iter_rows(values_only=True):
        






filename = data.xlsx